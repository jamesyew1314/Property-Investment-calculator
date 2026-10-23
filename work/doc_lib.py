from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_LINE_SPACING

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x5F, 0x69)
LIGHTGREY = RGBColor(0x8A, 0x90, 0x9A)
RED = RGBColor(0xA0, 0x32, 0x32)
GREEN = RGBColor(0x16, 0x78, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

EA_FONT = "PingFang SC"
LATIN_FONT = "Calibri"
MONO = "Menlo"

def _separate(doc):
    body = doc.element.body
    # python-docx keeps the section properties as the last body child; content sits before it
    idx = -1
    if len(body) and body[-1].tag == qn('w:sectPr'):
        idx = -2
    if len(body) >= abs(idx) and body[idx].tag == qn('w:tbl'):
        para(doc, "", size=2, space_after=2, space_before=0)

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        node = OxmlElement(f'w:{tag}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="BFC9D9", sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trPr.append(el)

def set_col_widths(table, widths_cm):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    for i, w in enumerate(widths_cm):
        table.columns[i].width = Cm(w)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)

def run_font(run, size=10.5, bold=False, color=DARK, italic=False, ea=EA_FONT, latin=LATIN_FONT):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)
    rFonts.set(qn('w:eastAsia'), ea)
    rFonts.set(qn('w:cs'), latin)

def para(doc, text="", size=10.5, bold=False, color=DARK, align=None, space_after=6,
         space_before=0, line=1.25, indent_cm=0, italic=False, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if indent_cm:
        pf.left_indent = Cm(indent_cm)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == 'justify':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if keep_next:
        pf.keep_with_next = True
    if text:
        r = p.add_run(text)
        run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p

def rich_para(doc, parts, size=10.5, space_after=6, line=1.25, align=None, indent_cm=0):
    """parts: list of (text, bold, color)"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent_cm:
        pf.left_indent = Cm(indent_cm)
    if align == 'justify':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, bold, color in parts:
        r = p.add_run(text)
        run_font(r, size=size, bold=bold, color=color)
    return p

def bullet(doc, text, level=0, size=10.5, color=DARK, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.2
    if bold_lead:
        r = p.add_run(bold_lead)
        run_font(r, size=size, bold=True, color=NAVY)
    r = p.add_run(text)
    run_font(r, size=size, color=color)
    return p

def numbered(doc, text, size=10.5, bold_lead=None):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if bold_lead:
        r = p.add_run(bold_lead)
        run_font(r, size=size, bold=True, color=NAVY)
    r = p.add_run(text)
    run_font(r, size=size)
    return p

def h1(doc, number, title, subtitle=None):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}")
    run_font(r, size=30, bold=True, color=RGBColor(0xC9,0xD4,0xE8))
    r2 = p.add_run(f"  {title}")
    run_font(r2, size=19, bold=True, color=NAVY)
    # rule
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12'); bottom.set(qn('w:space'), '6'); bottom.set(qn('w:color'), '1F3864')
    pBdr.append(bottom); pPr.append(pBdr)
    if subtitle:
        para(doc, subtitle, size=10, color=GREY, space_after=12, italic=True, keep_next=True)
    else:
        para(doc, "", size=4, space_after=8)

def h2(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    if number:
        r = p.add_run(f"{number}  ")
        run_font(r, size=13.5, bold=True, color=BLUE)
    r = p.add_run(title)
    run_font(r, size=13.5, bold=True, color=NAVY)

def h3(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    run_font(r, size=11, bold=True, color=RGBColor(0x33,0x44,0x55))

def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    run_font(r, size=8.5, color=LIGHTGREY, italic=True)

def table(doc, headers, rows, widths, font_size=8.8, header_size=8.8,
          align=None, caption_text=None, zebra=True, first_col_bold=False,
          header_fill="1F3864", highlight_rows=None, keep_rows=True):
    for _r in rows:
        if len(_r) != len(headers):
            raise ValueError(f"table row/header mismatch: {len(headers)} headers, {len(_r)} cells: {_r}")
    if caption_text:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(caption_text)
        run_font(r, size=9, bold=True, color=NAVY)
    _separate(doc)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)
    set_col_widths(t, widths)
    hdr = t.rows[0]
    set_repeat_header(hdr)
    set_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        run_font(r, size=header_size, bold=True, color=WHITE)
    highlight_rows = highlight_rows or []
    for ridx, row in enumerate(rows):
        tr = t.add_row()
        if keep_rows:
            set_cant_split(tr)
        for i, val in enumerate(row):
            cell = tr.cells[i]
            cell.text = ""
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if zebra and ridx % 2 == 1:
                set_cell_shading(cell, "F4F7FB")
            if ridx in highlight_rows:
                set_cell_shading(cell, "FFF1D9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if align and i < len(align) and align[i] == 'r':
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif align and i < len(align) and align[i] == 'c':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            run_font(r, size=font_size, bold=(first_col_bold and i == 0),
                     color=DARK)
    if caption_text:
        pass
    return t

def callout(doc, title, body, kind="info"):
    colors = {"info": ("EAF0FA", "2E5496"), "warn": ("FFF4E0", "C46E1E"), "risk": ("FBEAEA", "A03232"), "ok": ("E9F5EF", "16785A")}
    fill, border = colors.get(kind, colors["info"])
    _separate(doc)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(t, [17.0])
    set_cant_split(t.rows[0])
    cell = t.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '24' if edge == 'left' else '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), border)
        borders.append(el)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    run_font(r, size=10, bold=True, color=RGBColor.from_string(border))
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    run_font(r2, size=9.3, color=DARK)
    para(doc, "", size=4, space_after=4)

def code_block(doc, text, size=8.2):
    _separate(doc)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(t, [17.0])
    set_cant_split(t.rows[0])
    cell = t.rows[0].cells[0]
    set_cell_shading(cell, "F5F6F8")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    for idx, line in enumerate(text.split("\n")):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(line)
        run_font(r, size=size, color=RGBColor(0x33,0x33,0x33), latin=MONO, ea=MONO)
    para(doc, "", size=4, space_after=4)

def add_page_number_footer(section, left_text):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # tab stops: left text, right page
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), '9350'); tabs.append(tab)
    pPr.append(tabs)
    r = p.add_run(left_text)
    run_font(r, size=7.5, color=LIGHTGREY)
    r = p.add_run("\t")
    run_font(r, size=7.5, color=LIGHTGREY)
    r = p.add_run("Page ")
    run_font(r, size=7.5, color=LIGHTGREY)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    rr = OxmlElement('w:r'); rPr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '15'); rPr.append(sz)
    t = OxmlElement('w:t'); t.text = "1"; rr.append(rPr); rr.append(t); fld.append(rr)
    p._p.append(fld)
    r = p.add_run(" of ")
    run_font(r, size=7.5, color=LIGHTGREY)
    fld2 = OxmlElement('w:fldSimple'); fld2.set(qn('w:instr'), 'NUMPAGES')
    rr2 = OxmlElement('w:r'); rPr2 = OxmlElement('w:rPr'); sz2 = OxmlElement('w:sz'); sz2.set(qn('w:val'), '15'); rPr2.append(sz2)
    t2 = OxmlElement('w:t'); t2.text = "1"; rr2.append(rPr2); rr2.append(t2); fld2.append(rr2)
    p._p.append(fld2)

def add_header(section, text):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    r = p.add_run(text)
    run_font(r, size=7.5, color=LIGHTGREY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '4'); bottom.set(qn('w:space'), '4'); bottom.set(qn('w:color'), 'BFC9D9')
    pBdr.append(bottom); pPr.append(pBdr)

def setup_doc():
    doc = Document()
    # page
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    # normal style
    style = doc.styles['Normal']
    style.font.name = LATIN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), EA_FONT)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.25
    # list styles font
    for sname in ('List Bullet', 'List Bullet 2', 'List Number'):
        st = doc.styles[sname]
        st.font.name = LATIN_FONT
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), EA_FONT)
    section.different_first_page_header_footer = True
    add_header(section, "Property Investment Analysis System  ·  需求深度分析与开发方案 v1.0")
    add_page_number_footer(section, "Confidential  ·  Prepared for Andy Lau / EAG  ·  2026-10-18")
    return doc

def cover(doc, title, subtitle, meta_lines):
    # top spacing
    for _ in range(3):
        para(doc, "", size=10, space_after=0)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("PROPERTY INVESTMENT ANALYSIS SYSTEM")
    run_font(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    run_font(r, size=34, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(subtitle)
    run_font(r, size=15, color=GREY)
    # rule
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '18'); bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1F3864')
    pBdr.append(bottom); pPr.append(pBdr)
    para(doc, "", size=10, space_after=18)
    for k, v in meta_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(k)
        run_font(r, size=9, color=LIGHTGREY)
        r = p.add_run("   " + v)
        run_font(r, size=10, bold=True, color=DARK)
    para(doc, "", size=10, space_after=10)
    callout(doc, "阅读提示",
            "这份文件分两部分：第 2–4 章是“深度分析”（我从你的 Excel、两份报价单和截图里读到了什么、发现了什么问题）；第 5–15 章是“方案”（数据模型、公式、界面、架构、路线图、验收标准）。程序师可以直接按第 7–13 章报价与开发。",
            kind="info")
