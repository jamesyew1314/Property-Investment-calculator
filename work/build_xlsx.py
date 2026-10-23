from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from datetime import date

NAVY="1F3864"; BLUE="2E5496"; LIGHT="EAF0FA"; INPUT="FFF2CC"; GREY="595F69"; ZEBRA="F4F7FB"
thin=Side(style="thin", color="BFC9D9")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

def title(ws, cell, text):
    ws[cell]=text
    ws[cell].font=Font(name="Calibri", size=18, bold=True, color=NAVY)
    ws.row_dimensions[int(''.join(c for c in cell if c.isdigit()))].height=26

def section(ws, row, text, ncols=8):
    c=ws.cell(row=row, column=2, value=text)
    c.font=Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col in range(2, 2+ncols):
        ws.cell(row=row, column=col).fill=PatternFill("solid", fgColor=BLUE)
        ws.cell(row=row, column=col).border=border
    ws.row_dimensions[row].height=20

def label(ws, row, col, text, bold=True):
    c=ws.cell(row=row, column=col, value=text)
    c.font=Font(name="Calibri", size=10, bold=bold, color="222222")
    return c

def inp(ws, row, col, value, fmt=None):
    c=ws.cell(row=row, column=col, value=value)
    c.fill=PatternFill("solid", fgColor=INPUT)
    c.border=border
    c.font=Font(name="Calibri", size=10, color="222222")
    if fmt: c.number_format=fmt
    return c

def out(ws, row, col, value, fmt=None, bold=False):
    c=ws.cell(row=row, column=col, value=value)
    c.fill=PatternFill("solid", fgColor=LIGHT)
    c.border=border
    c.font=Font(name="Calibri", size=10, bold=bold, color="222222")
    if fmt: c.number_format=fmt
    return c

def note(ws, row, text, ncols=8):
    c=ws.cell(row=row, column=2, value=text)
    c.font=Font(name="Calibri", size=9, italic=True, color=GREY)
    c.alignment=Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=1+ncols)
    ws.row_dimensions[row].height=28

def widths(ws, d):
    for col,w in d.items():
        ws.column_dimensions[col].width=w

wb=Workbook()

# ---------------- Read Me ----------------
ws=wb.active; ws.title="Read Me"
widths(ws, {"A":2,"B":26,"C":100})
title(ws,"B2","投资系统 — 公式验证与案例对账 v1.1")
ws["B4"]="用途"; ws["C4"]="验证法定费用公式、4 个 Excel 案例的现金流口径，以及贷款摊销/XIRR 的实现。"
ws["B5"]="来源"; ws["C5"]="Andy Lau - Invest Formula (Long Term) Year2026 (4).xlsx；SPA Quotation_20260526.pdf；Loan Quotation_20260526.pdf。"
ws["B6"]="使用"; ws["C6"]="黄色单元格可改；蓝色单元格是公式结果；灰色斜体是说明。"
ws["B7"]="口径"; ws["C7"]="金额 RM；ROI = 净利润 ÷ 现金投入（Cash-on-Cash），与 Excel 一致；复核口径见各案例页。"
ws["B8"]="免责"; ws["C8"]="法定费率以报价单/律师确认为准；RPGT/SST 为估算，不构成税务意见。"
for r in range(4,9):
    ws.cell(row=r,column=2).font=Font(name="Calibri", size=10, bold=True, color=NAVY)
    ws.cell(row=r,column=2).alignment=Alignment(vertical="top")
    ws.cell(row=r,column=3).font=Font(name="Calibri", size=10, color="222222")
    ws.cell(row=r,column=3).alignment=Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height=32

ws["B11"]="目录"; ws["B11"].font=Font(size=12,bold=True,color=NAVY)
rows=[("Summary","4 个案例：Excel vs 复核（ROI/净利润/关键差异）"),
      ("Fee Calculator","法定费用：S&P/贷款律师费、印花税、SST；与报价单对账"),
      ("Case A","Austin Perdana：SPA 1.6M / 实价 1.25M / 贷款 1.28M"),
      ("Case B","L.T：SPA 900k / 实价 800k / 贷款 720k / 25–27 个月"),
      ("Case C","Cashback：SPA 900k / 实价 735k / 贷款 765k"),
      ("Case D","Downpay：SPA 900k / 实价 800k / 贷款 720k"),
      ("Loan Amortization","720k / 4.5% / 30 年；含 24 个月余额"),
      ("Cash Flow & XIRR","Case C/D 的日期化现金流与 XIRR"),
      ("KPI Definitions","指标定义、分母、常见错误")]
r=12
for a,b in rows:
    ws.cell(row=r,column=2,value=a).font=Font(size=10,bold=True,color=BLUE)
    ws.cell(row=r,column=3,value=b).font=Font(size=10,color="222222")
    r+=1

# ---------------- Summary ----------------
ws=wb.create_sheet("Summary",1)
widths(ws, {"A":2,"B":22,"C":10,"D":12,"E":12,"F":16,"G":16,"H":60})
title(ws,"B2","4 个案例 — Excel vs 复核")
section(ws,4,"Summary",7)
headers=["案例","持有月","Excel ROI","复核 ROI","Excel 净利润","复核净利润","关键差异"]
for i,h in enumerate(headers):
    c=ws.cell(row=5,column=2+i,value=h); c.font=Font(size=9,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.border=border
    c.alignment=Alignment(wrap_text=True,horizontal="center")
rows=[("A  Austin Perdana",17,1.56318806748339,1.58709111712041,187419.34,190285.21,"SPA 1.60M vs 实价 1.25M；赎回 > 贷款；缺结算单"),
      ("B  Profit Calculator (L.T)","25–27",0.09112489353948058,0.0455819479555238,16472.76,8016.04,"漏出售税 3,545；本金 28,349 vs 23,438"),
      ("C  Loan Cashback",13,1.14463549132331,1.76165753328209,48040.26,48040.26,"定金 14,700 重复；现金投入 33,970 vs 19,270"),
      ("D  Loan Downpay",23,0.186580593294365,0.202312589746739,36687.76,36507.76,"定金 16,000 重复；本金差 180")]
r=6
for row in rows:
    for i,v in enumerate(row):
        c=ws.cell(row=r,column=2+i,value=v)
        c.font=Font(size=9.5,color="222222")
        c.border=border
        c.alignment=Alignment(vertical="center", wrap_text=(i==6))
        if i in (2,3): c.number_format="0.00%"
        if i in (4,5): c.number_format="#,##0.00"
        if r%2==1: c.fill=PatternFill("solid",fgColor=ZEBRA)
    ws.row_dimensions[r].height=40
    r+=1
note(ws,11,"复核口径：以银行赎回余额统一本金；定金只算一次；出售税进利润；预收租金/押金按负债处理。详见各案例页。",7)
ws.freeze_panes="A6"

# ---------------- Fee Calculator ----------------
ws=wb.create_sheet("Fee Calculator")
widths(ws, {"A":2,"B":30,"C":18,"D":18,"E":18,"F":18,"G":18,"H":18})
title(ws,"B2","法定费用计算器")
ws["B4"]="输入"; ws["B4"].font=Font(size=11,bold=True,color="FFFFFF"); ws["B4"].fill=PatternFill("solid",fgColor=BLUE)
for col in "CDEFGH":
    ws[f"{col}4"].fill=PatternFill("solid",fgColor=BLUE)

label(ws,5,2,"SPA Price / Purchase Price (RM)"); inp(ws,5,3,1600000,"#,##0.00")
label(ws,6,2,"Loan Amount (RM)"); inp(ws,6,3,1280000,"#,##0.00")
label(ws,7,2,"SST rate"); inp(ws,7,3,0.08,"0.00%")
label(ws,8,2,"Quote total (optional)"); inp(ws,8,3,77487,"#,##0.00")

section(ws,10,"Outputs — statutory only",6)
outs=[("S&P legal fee","=0.0125*MIN(C5,500000)+0.01*MIN(MAX(C5-500000,0),7000000)+0.01*MAX(C5-7500000,0)"),
      ("S&P stamp duty","=0.01*MIN(C5,100000)+0.02*MIN(MAX(C5-100000,0),400000)+0.03*MIN(MAX(C5-500000,0),500000)+0.04*MAX(C5-1000000,0)"),
      ("S&P SST","=C7*C11"),
      ("S&P statutory total","=SUM(C11:C13)"),
      ("Loan legal fee","=0.0125*MIN(C6,500000)+0.01*MIN(MAX(C6-500000,0),7000000)+0.01*MAX(C6-7500000,0)"),
      ("Loan stamp duty","=0.005*C6"),
      ("Loan SST","=C7*C15"),
      ("Loan statutory total","=SUM(C15:C17)"),
      ("Combined statutory","=C14+C18"),
      ("Quote − S&P statutory","=C8-C14"),
      ("Quote − combined","=C8-C19")]
r=11
for lab,formula in outs:
    label(ws,r,2,lab); out(ws,r,3,formula,"#,##0.00", bold=("total" in lab or "Combined" in lab)); r+=1

note(ws,23,"报价单的 Total 含 SST、贷款 stamping、律所 disbursement；不能把 Total 当成“律师费+印花税”。S&P 报价可能不含贷款律师费，所以 Quote − combined 为负是正常的。",6)

sp_sec=25
section(ws,sp_sec,"S&P 报价单对账（PDF 第 1–5 页）",7)
headers=["Purchase Price","S&P legal fee","S&P stamp duty","Statutory total","Quote Total 1","Residual","Residual / Price"]
for i,h in enumerate(headers):
    c=ws.cell(row=sp_sec+1,column=2+i,value=h); c.font=Font(size=9,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.border=border
    c.alignment=Alignment(wrap_text=True,horizontal="center")
prices=[50000,100000,240000,500000,900000,1000000,1600000,2000000]
quotes=[4612,5857,10997,21107,39427,44007,77487,99807]
r=sp_sec+2
for p,q in zip(prices,quotes):
    inp(ws,r,2,p,"#,##0")
    ws.cell(row=r,column=3,value="=0.0125*MIN(B%d,500000)+0.01*MIN(MAX(B%d-500000,0),7000000)+0.01*MAX(B%d-7500000,0)"%(r,r,r)).number_format="#,##0.00"
    ws.cell(row=r,column=4,value="=0.01*MIN(B%d,100000)+0.02*MIN(MAX(B%d-100000,0),400000)+0.03*MIN(MAX(B%d-500000,0),500000)+0.04*MAX(B%d-1000000,0)"%(r,r,r,r)).number_format="#,##0.00"
    ws.cell(row=r,column=5,value=f"=C{r}+D{r}").number_format="#,##0.00"
    inp(ws,r,6,q,"#,##0.00")
    ws.cell(row=r,column=7,value=f"=F{r}-E{r}").number_format="#,##0.00"
    ws.cell(row=r,column=8,value=f"=G{r}/B{r}").number_format="0.00%"
    for col in range(2,9):
        ws.cell(row=r,column=col).border=border
    r+=1
note(ws,r,"Residual 从 3,487 到 14,557，说明它包含固定杂费 + 按比例收费；报价单未拆分。",7)

loan_sec=r+2
section(ws,loan_sec,"Loan 报价单对账（PDF 第 1–2 页）",6)
headers=["Loan Amount","Loan legal fee","Loan stamp duty","Statutory total","Quote Total 1","Residual"]
for i,h in enumerate(headers):
    c=ws.cell(row=loan_sec+1,column=2+i,value=h); c.font=Font(size=9,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.border=border
    c.alignment=Alignment(wrap_text=True,horizontal="center")
loans=[50000,100000,300000]
lquotes=[5705,6630,10330]
rr=loan_sec+2
for p,q in zip(loans,lquotes):
    inp(ws,rr,2,p,"#,##0")
    ws.cell(row=rr,column=3,value="=0.0125*MIN(B%d,500000)+0.01*MIN(MAX(B%d-500000,0),7000000)+0.01*MAX(B%d-7500000,0)"%(rr,rr,rr)).number_format="#,##0.00"
    ws.cell(row=rr,column=4,value=f"=0.005*B{rr}").number_format="#,##0.00"
    ws.cell(row=rr,column=5,value=f"=C{rr}+D{rr}").number_format="#,##0.00"
    inp(ws,rr,6,q,"#,##0.00")
    ws.cell(row=rr,column=7,value=f"=F{rr}-E{rr}").number_format="#,##0.00"
    for col in range(2,8):
        ws.cell(row=rr,column=col).border=border
    rr+=1

# ---------------- Case sheets ----------------
def case_sheet(name, title_text, inputs, calcs, notes, checks):
    """inputs: list of (key, label, value, fmt)
       calcs: list of (label, formula_template, fmt, bold)
       checks: list of (label, formula_template, fmt)
    """
    ws=wb.create_sheet(name)
    widths(ws, {"A":2,"B":36,"C":18,"D":18,"E":18,"F":46})
    title(ws,"B2",title_text)
    section(ws,4,"Inputs",4)
    addr={}
    r=5
    for key,lab,val,fmt in inputs:
        label(ws,r,2,lab); inp(ws,r,3,val,fmt); addr[key]=f"C{r}"; r+=1
    section(ws,r+1,"Calculations",4)
    r+=2
    for lab,formula,fmt,bold in calcs:
        label(ws,r,2,lab)
        f=formula.format(**addr) if isinstance(formula,str) and "{" in formula else formula
        out(ws,r,3,f,fmt,bold); addr[lab]=f"C{r}"; r+=1
    section(ws,r+1,"Checks / Notes",4)
    r+=2
    for lab,formula,fmt in checks:
        label(ws,r,2,lab)
        f=formula.format(**addr) if isinstance(formula,str) and "{" in formula else formula
        out(ws,r,3,f,fmt); addr[lab]=f"C{r}"; r+=1
    for n in notes:
        note(ws,r,n,5); r+=1
    ws.freeze_panes="A5"
    return ws

case_sheet("Case A","Case A — Austin Perdana（部分复核）",
    [("price","买入价",1250000,"#,##0.00"),("spa","SPA 价",1600000,"#,##0.00"),
     ("loan","贷款",1280000,"#,##0.00"),("deposit","定金",100000,"#,##0.00"),
     ("costs","购买成本（含 RPGT/MRTA）",152895.58,"#,##0.00"),
     ("cash_excel","实际现金投入（Excel）",119895.58,"#,##0.00"),
     ("sale","出售价",1620000,"#,##0.00"),("redemption","赎回",1296950,"#,##0.00"),
     ("selling","出售费用",41669.21,"#,##0.00"),("retention","RPGT retention",48600,"#,##0.00"),
     ("deposit_refund","押金退还",12300,"#,##0.00"),("advance_rent","预收租金退还",2865.87,"#,##0.00")],
    [("Total Invest", "={price}+{costs}", "#,##0.00", True),
     ("出售现金回报", "={sale}-{redemption}-{selling}-{retention}-{deposit_refund}-{advance_rent}", "#,##0.00", True),
     ("Excel 净利润", 187419.34, "#,##0.00", True),
     ("复核净利润（剔除预收租金）", "={Excel 净利润}+{advance_rent}", "#,##0.00", True),
     ("Excel ROI", "={Excel 净利润}/{cash_excel}", "0.00%", True),
     ("复核 ROI", "={复核净利润（剔除预收租金）}/{cash_excel}", "0.00%", True)],
    ["Case A 的 SPA 1.60M 与实价 1.25M 差额 350,000 缺结算单；赎回 1,296,950 > 贷款 1,280,000，需银行 settlement。",
     "RPGT retention 48,600 是现金流暂扣，不是最终税额；30,000 被放进购买成本，性质待确认。",
     "本页只做部分复核；不要用 ROI 做最终决策，直到补齐 completion statement。"],
    [("现金回报 − 现金投入 − 净利润（应为 0）", "={出售现金回报}-{cash_excel}-{Excel 净利润}", "#,##0.00"),
     ("现金回报 − 现金投入 − 复核净利润（应为 0）", "={出售现金回报}-{cash_excel}-{复核净利润（剔除预收租金）}", "#,##0.00")])

case_sheet("Case B","Case B — Profit Calculator (L.T)",
    [("price","实价",800000,"#,##0.00"),("spa","SPA 价",900000,"#,##0.00"),
     ("loan","贷款",720000,"#,##0.00"),("deposit","定金",16000,"#,##0.00"),
     ("costs","购买成本",72422.24,"#,##0.00"),("rent","租金",72200,"#,##0.00"),
     ("instalment","供款",100000,"#,##0.00"),("principal_excel","已还本金（Excel）",28349,"#,##0.00"),
     ("sale","出售价",910000,"#,##0.00"),("redemption","赎回",696562.27,"#,##0.00"),
     ("selling","出售费用",25199.45,"#,##0.00")],
    [("尾款", "={price}-{loan}-{deposit}", "#,##0.00", False),
     ("初始现金", "={deposit}+{costs}+{尾款}", "#,##0.00", True),
     ("月度净现金流（25 个月）", "={rent}-{instalment}", "#,##0.00", False),
     ("净售价", "={sale}-{redemption}-{selling}", "#,##0.00", True),
     ("净利润（复核）", "={净售价}+{月度净现金流（25 个月）}-{初始现金}", "#,##0.00", True),
     ("已还本金（余额反推）", "={loan}-{redemption}", "#,##0.00", False),
     ("ROI（复核）", "={净利润（复核）}/({初始现金}+{已还本金（余额反推）})", "0.00%", True),
     ("Excel ROI", 0.09112489353948058, "0.00%", False)],
    ["Excel 的 gross profit 漏了 3,545.45 的出售税；本页用净售价（已扣税）重算。",
     "Excel 用 28,349 作为本金，但赎回余额反推 23,437.73；本页用银行余额。",
     "Excel 的 25 个月与 27 个月不同；ROI 对比时要注意持有期。"],
    [("Excel 净利润", 16472.76, "#,##0.00"),
     ("差异（Excel − 复核）", "={Excel 净利润}-{净利润（复核）}", "#,##0.00"),
     ("利息（供款 − 本金）", "={instalment}-{已还本金（余额反推）}", "#,##0.00")])

case_sheet("Case C","Case C — Loan Cashback（定金重复计算）",
    [("price","实价",735000,"#,##0.00"),("spa","SPA 价",900000,"#,##0.00"),
     ("loan","贷款",765000,"#,##0.00"),("deposit","定金",14700,"#,##0.00"),
     ("costs","购买成本",49269.92,"#,##0.00"),("rent","租金",24700,"#,##0.00"),
     ("instalment","供款",50389.82,"#,##0.00"),("principal","已还本金",8000,"#,##0.00"),
     ("sale","出售价",850000,"#,##0.00"),("redemption","赎回",757000,"#,##0.00")],
    [("交割现金流", "={loan}-({price}-{deposit})-{costs}", "#,##0.00", False),
     ("初始现金（复核）", "={deposit}+MAX(0,-{交割现金流})", "#,##0.00", True),
     ("Excel 初始现金", "={deposit}+{costs}-({loan}-{price})", "#,##0.00", False),
     ("月度净现金流（13 个月）", "={rent}-{instalment}", "#,##0.00", False),
     ("净售价", "={sale}-{redemption}", "#,##0.00", True),
     ("净利润（复核）", "={净售价}+{月度净现金流（13 个月）}-{初始现金（复核）}", "#,##0.00", True),
     ("ROI（复核，含本金）", "={净利润（复核）}/({初始现金（复核）}+{principal})", "0.00%", True),
     ("Excel ROI", 1.1446354913233088, "0.00%", False)],
    ["Excel 的 D16 = 735,000 − 765,000 = −30,000，没有扣定金；D18 又加了 14,700。",
     "正确：尾款 = 735,000 − 14,700 = 720,300；交割现金流 = 765,000 − 720,300 − 49,269.92 = −4,569.92。",
     "净利润 48,040.26 是对的；错的是现金投入 33,969.92 vs 19,269.92。"],
    [("初始现金差（Excel − 复核）", "={Excel 初始现金}-{初始现金（复核）}", "#,##0.00"),
     ("净利润（Excel）", 48040.26, "#,##0.00"),
     ("ROI 差（Excel − 复核）", "={Excel ROI}-{ROI（复核，含本金）}", "0.00%")])

case_sheet("Case D","Case D — Loan Downpay（定金重复计算）",
    [("price","实价",800000,"#,##0.00"),("spa","SPA 价",900000,"#,##0.00"),
     ("loan","贷款",720000,"#,##0.00"),("deposit","定金",16000,"#,##0.00"),
     ("costs","购买成本",72422.24,"#,##0.00"),("rent","租金",62900,"#,##0.00"),
     ("instalment","供款",92000,"#,##0.00"),("principal_excel","已还本金（Excel）",28210,"#,##0.00"),
     ("sale","出售价",910000,"#,##0.00"),("redemption","赎回",691970,"#,##0.00")],
    [("尾款", "={price}-{loan}-{deposit}", "#,##0.00", False),
     ("初始现金（复核）", "={deposit}+{costs}+{尾款}", "#,##0.00", True),
     ("Excel 初始现金", "={deposit}+{costs}+({price}-{loan})", "#,##0.00", False),
     ("月度净现金流（23 个月）", "={rent}-{instalment}", "#,##0.00", False),
     ("净售价", "={sale}-{redemption}", "#,##0.00", True),
     ("净利润（复核）", "={净售价}+{月度净现金流（23 个月）}-{初始现金（复核）}", "#,##0.00", True),
     ("已还本金（余额反推）", "={loan}-{redemption}", "#,##0.00", False),
     ("ROI（复核）", "={净利润（复核）}/({初始现金（复核）}+{已还本金（余额反推）})", "0.00%", True),
     ("Excel ROI", 0.18658059329436463, "0.00%", False)],
    ["Excel D17 = 800,000 − 720,000 = 80,000，D19 又加定金 16,000；定金算了两次。",
     "正确尾款 = 64,000；初始现金 = 152,422.24，不是 168,422.24。",
     "本金 28,210 与赎回余额反推 28,030 差 180，导致净利润 36,687.76 vs 36,507.76。"],
    [("初始现金差（Excel − 复核）", "={Excel 初始现金}-{初始现金（复核）}", "#,##0.00"),
     ("净利润（Excel）", 36687.76, "#,##0.00"),
     ("ROI 差（Excel − 复核）", "={Excel ROI}-{ROI（复核）}", "0.00%")])

# ---------------- Loan Amortization ----------------
ws=wb.create_sheet("Loan Amortization")
widths(ws, {"A":2,"B":14,"C":14,"D":16,"E":16,"F":16,"G":16,"H":16})
title(ws,"B2","贷款摊销 — 720k / 4.5% / 30 年")
label(ws,4,2,"Loan"); inp(ws,4,3,720000,"#,##0.00")
label(ws,5,2,"Annual rate"); inp(ws,5,3,0.045,"0.00%")
label(ws,6,2,"Years"); inp(ws,6,3,30,"0")
label(ws,7,2,"Start date"); inp(ws,7,3,"2023-01-01","dd/mm/yyyy")
label(ws,8,2,"Monthly payment"); out(ws,8,3,"=PMT(C5/12,C6*12,-C4)","#,##0.00", True)
label(ws,9,2,"Balance at month 24"); out(ws,9,3,"=INDEX(H13:H372,24)","#,##0.00", True)

headers=["Month","Date","Opening","Payment","Interest","Principal","Closing"]
for i,h in enumerate(headers):
    c=ws.cell(row=12,column=2+i,value=h); c.font=Font(size=9,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.border=border
    c.alignment=Alignment(horizontal="center")
for m in range(1,361):
    r=12+m
    ws.cell(row=r,column=2,value=m).border=border
    ws.cell(row=r,column=3,value=f"=EDATE($C$7,B{r})").number_format="dd/mm/yyyy"
    ws.cell(row=r,column=4,value=("=$C$4" if m==1 else f"=H{r-1}")).number_format="#,##0.00"
    ws.cell(row=r,column=5,value="=$C$8").number_format="#,##0.00"
    ws.cell(row=r,column=6,value=f"=D{r}*$C$5/12").number_format="#,##0.00"
    ws.cell(row=r,column=7,value=f"=E{r}-F{r}").number_format="#,##0.00"
    ws.cell(row=r,column=8,value=f"=D{r}-G{r}").number_format="#,##0.00"
    for col in range(2,9):
        ws.cell(row=r,column=col).border=border
    if m%2==0:
        for col in range(2,9):
            ws.cell(row=r,column=col).fill=PatternFill("solid",fgColor=ZEBRA)
note(ws,375,"这是 monthly rest 的简化表；daily rest 需按日计息。银行 settlement 优先于本表。",7)

# ---------------- Cash Flow & XIRR ----------------
ws=wb.create_sheet("Cash Flow & XIRR")
widths(ws, {"A":2,"B":14,"C":16,"D":18,"E":60})
title(ws,"B2","Case C — 日期化现金流与 XIRR")
for i,h in enumerate(["Date","Amount","Label","Note"]):
    c=ws.cell(row=4,column=2+i,value=h); c.font=Font(size=10,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.border=border
flows=[(date(2023,10,1),-14700,"Deposit","定金"),
       (date(2023,11,1),-4569.92,"Completion","交割现金流"),
       ]
r=5
for d,a,l,n in flows:
    ws.cell(row=r,column=2,value=d).number_format="dd/mm/yyyy"
    ws.cell(row=r,column=3,value=a).number_format="#,##0.00"
    ws.cell(row=r,column=4,value=l); ws.cell(row=r,column=5,value=n)
    r+=1
# 13 months rent - installment
for m in range(1,14):
    ws.cell(row=r,column=2,value=f"=EDATE(DATE(2023,11,1),{m})").number_format="dd/mm/yyyy"
    ws.cell(row=r,column=3,value=1900-3876.14).number_format="#,##0.00"
    ws.cell(row=r,column=4,value="Monthly net").font=Font(size=9,color=GREY)
    r+=1
ws.cell(row=r,column=2,value=date(2024,12,31)).number_format="dd/mm/yyyy"
ws.cell(row=r,column=3,value=850000-757000).number_format="#,##0.00"
ws.cell(row=r,column=4,value="Net sale")
last=r
label(ws,last+2,2,"Sum of flows"); out(ws,last+2,3,f"=SUM(C5:C{last})","#,##0.00",True)
label(ws,last+3,2,"XIRR"); out(ws,last+3,3,f"=XIRR(C5:C{last},B5:B{last})","0.00%",True)
note(ws,last+5,"日期为假设；真实 XIRR 需按实际付款日。XIRR 对日期敏感，差几天可能差几个百分点。",4)

# ---------------- KPI Definitions ----------------
ws=wb.create_sheet("KPI Definitions")
widths(ws, {"A":2,"B":30,"C":55,"D":32,"E":40})
title(ws,"B2","指标定义与分母")
headers=["KPI","Definition / Formula","Denominator","Common error"]
for i,h in enumerate(headers):
    c=ws.cell(row=4,column=2+i,value=h); c.font=Font(size=10,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.border=border
kpis=[("Total Acquisition Cost","actual price + legal/stamp/valuation/agent/consent","excl. renovation/holding","including renovation"),
      ("Total Project Cost","acquisition + renovation + holding + interest + selling","all-in","missing interest"),
      ("Initial Cash Out","deposit + own-fund costs + completion shortfall − cashback","own cash","loan-funded costs counted"),
      ("Completion Cash","loan + rebate − (price − deposit) − costs at completion","completion","deposit double count"),
      ("Monthly Net Cash Flow","effective rent + other income − opex − instalment − tax − insurance − repair","monthly","rent − instalment only"),
      ("Peak Cash Requirement","max(0, −min(cumulative)) + reserve","max funding","initial cash only"),
      ("Principal Repaid","loan principal − outstanding","cumulative","payment × months"),
      ("Interest Paid","Σ interest = Σ payment − principal","cumulative","principal as expense"),
      ("Net Sale Proceeds","price − redemption − selling costs − retention − deposits − advance rent","completion","retention as permanent cost"),
      ("Net Profit (cash)","net sale + cumulative net cash − total cash invested","cash","cash return as profit"),
      ("Total ROI","net profit / total project cost","total cost","cash denominator"),
      ("Cash-on-Cash","net profit / total cash invested","cash","denominator undefined"),
      ("Simple Annualized","total ROI × 365 / days","annual","hand-filled months"),
      ("CAGR","(ending/beginning)^(365/days) − 1","compound","CAGR vs XIRR"),
      ("XIRR","solve Σ CF/(1+r)^(days/365)=0","dated flows","missing cash flow"),
      ("Gross Yield","annual rent / price","annual","rent/loan"),
      ("Net Yield","(rent − opex) / total cost","annual","gross rent"),
      ("Break-even Price","solve net sale + rent − cash invested = 0","price","ignore costs"),
      ("DSCR","NOI / annual debt service","coverage","gross rent"),
      ("RPGT","rate × max(0, gain − exemption)","tax","retention as tax")]
r=5
for row in kpis:
    for i,v in enumerate(row):
        c=ws.cell(row=r,column=2+i,value=v)
        c.font=Font(size=9,color="222222")
        c.alignment=Alignment(wrap_text=True,vertical="top")
        c.border=border
        if r%2==0: c.fill=PatternFill("solid",fgColor=ZEBRA)
    ws.row_dimensions[r].height=30
    r+=1

for ws in wb.worksheets:
    ws.sheet_view.showGridLines=False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

wb.save("work/validation.xlsx")
print("saved")
